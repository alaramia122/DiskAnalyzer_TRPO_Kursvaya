#!/usr/bin/env python3
"""Fill the generated report's static table of contents with rendered page numbers."""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Cm
except ImportError as exc:
    raise SystemExit("Missing dependency: python-docx") from exc


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def extract_pages(pdf_path: Path) -> list[str]:
    result = subprocess.run(
        ["pdftotext", "-layout", "-enc", "UTF-8", str(pdf_path), "-"],
        check=True, capture_output=True, text=True, encoding="utf-8",
    )
    return result.stdout.split("\f")


def locate_heading(title: str, pages: list[str]) -> int:
    target = normalize(title)
    # The first page is the title sheet and the second is the table of contents.
    # Search from page 3 so that a TOC entry cannot be mistaken for its own target.
    for page_number, page in enumerate(pages, start=1):
        if page_number <= 2:
            continue
        lines = [normalize(line) for line in page.splitlines()]
        if target in lines:
            return page_number
    # Fallback for a heading wrapped by the renderer.
    for page_number, page in enumerate(pages, start=1):
        if page_number <= 2:
            continue
        if target in normalize(page):
            return page_number
    raise RuntimeError(f"Could not locate heading in rendered PDF: {title!r}")


def main() -> int:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: update_toc_page_numbers.py REPORT.docx REPORT.pdf")
    docx_path = Path(sys.argv[1]).resolve()
    pdf_path = Path(sys.argv[2]).resolve()
    if not docx_path.is_file() or not pdf_path.is_file():
        raise SystemExit("Both DOCX and rendered PDF must exist.")

    pages = extract_pages(pdf_path)
    doc = Document(docx_path)
    entries = [p for p in doc.paragraphs if p.style.name.startswith("TOC Entry ")]
    if not entries:
        raise RuntimeError("No marked TOC entries found in DOCX.")

    for paragraph in entries:
        title = paragraph.text.split("\t", 1)[0].strip()
        page_number = locate_heading(title, pages)
        for run in list(paragraph.runs):
            run._element.getparent().remove(run._element)
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        paragraph.add_run(title)
        paragraph.add_run("\t")
        page_run = paragraph.add_run(str(page_number))
        page_run.bold = False

    doc.save(docx_path)
    print(f"Updated TOC page numbers: {len(entries)} entries.")
    print("Page numbers were derived from the rendered PDF; render again to verify layout.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
